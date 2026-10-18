"""Generación de los libros Excel con openpyxl (Semana 17, HU-143).

Dos hojas, sin fórmulas: «Resumen» (título, filtros, generado por,
fecha de Lima, número de filas, notas y resumen) y «Datos» (una fila
por registro con los tipos reales: números como números y fechas como
fechas). Encabezado congelado y autofiltro en «Datos», anchos de
columna razonables y todo texto que empiece con =, +, - o @ escrito
como cadena inofensiva —nunca como fórmula—.

El libro se construye en memoria (BytesIO) y no se guarda en disco.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import catalogo, formato
from .consultas import fila_totales

_TITULO = Font(bold=True, size=14)
_CABECERA = Font(bold=True)


def generar_excel(reporte, datos, generado_por, generado_en):
    """Libro .xlsx en memoria: hoja «Resumen» y hoja «Datos»."""
    libro = Workbook()
    hoja_resumen = libro.active
    hoja_resumen.title = "Resumen"
    _escribir_resumen(hoja_resumen, reporte, datos, generado_por, generado_en)
    hoja_datos = libro.create_sheet("Datos")
    _escribir_datos(hoja_datos, reporte, datos["filas"])
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _celda(hoja, fila, columna, valor, negrita=False):
    celda = hoja.cell(row=fila, column=columna, value=valor)
    if negrita:
        celda.font = _CABECERA
    if formato.es_texto_de_formula(valor):
        # El texto riesgoso se guarda como cadena, jamás como fórmula
        celda.data_type = "s"
    return celda


def _escribir_resumen(hoja, reporte, datos, generado_por, generado_en):
    filas = [
        (reporte["nombre"], _TITULO),
        (f"Filtros: {datos['filtros_texto']}", None),
        (f"Generado por: {generado_por}", None),
        (
            f"Generado el: {formato.hora_larga(generado_en)} (hora de Lima)",
            None,
        ),
        (f"Filas: {len(datos['filas'])}", None),
        (catalogo.AVISO_PRIVACIDAD, None),
    ]
    fila = 1
    for texto_fila, fuente in filas:
        celda = hoja.cell(row=fila, column=1, value=texto_fila)
        if fuente is not None:
            celda.font = fuente
        fila += 1
    for advertencia in datos["advertencias"]:
        hoja.cell(row=fila, column=1, value=f"Nota: {advertencia}")
        fila += 1
    fila += 1  # separación

    resumen = datos["resumen"]
    if reporte["tipo"] == "productores":
        valores = (
            ("Productores distintos", resumen["productores"]),
            ("Filas", resumen["filas"]),
            (
                "Superficie de parcelas (ha)",
                resumen["superficie_parcelas_ha"],
            ),
            ("Superficie sembrada (ha)", resumen["superficie_sembrada_ha"]),
        )
        for etiqueta, valor in valores:
            _celda(hoja, fila, 1, etiqueta, negrita=True)
            _celda(hoja, fila, 2, valor)
            fila += 1
    else:
        columnas = catalogo.COLUMNAS_RESUMEN_PRODUCTIVO
        for indice, columna in enumerate(columnas, start=1):
            _celda(hoja, fila, indice, columna["etiqueta"], negrita=True)
        fila += 1
        for fila_datos in resumen["por_cultivo"]:
            valores = formato.celda_excel(columnas, fila_datos)
            for indice, columna in enumerate(columnas, start=1):
                _celda(hoja, fila, indice, valores[columna["clave"]])
            fila += 1
        totales = fila_totales(resumen)
        valores = formato.celda_excel(columnas, totales)
        for indice, columna in enumerate(columnas, start=1):
            _celda(hoja, fila, indice, valores[columna["clave"]], negrita=True)
    _ajustar_anchos(hoja)


def _escribir_datos(hoja, reporte, filas):
    columnas = reporte["columnas"]
    for indice, columna in enumerate(columnas, start=1):
        _celda(hoja, 1, indice, columna["etiqueta"], negrita=True)
    for numero_fila, datos_fila in enumerate(filas, start=2):
        valores = formato.celda_excel(columnas, datos_fila)
        for indice, columna in enumerate(columnas, start=1):
            _celda(hoja, numero_fila, indice, valores[columna["clave"]])
    hoja.freeze_panes = "A2"
    if filas:
        ultima = get_column_letter(len(columnas))
        hoja.auto_filter.ref = f"A1:{ultima}{len(filas) + 1}"
    _ajustar_anchos(hoja)


def _ajustar_anchos(hoja):
    """Anchos razonables (10..45) según el contenido de cada columna."""
    for indice in range(1, hoja.max_column + 1):
        ancho = 10
        for fila in range(1, min(hoja.max_row, 200) + 1):
            valor = hoja.cell(row=fila, column=indice).value
            if valor is None:
                continue
            ancho = max(ancho, min(45, len(str(valor)) + 2))
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
