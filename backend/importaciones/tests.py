"""Pruebas 326-345 (Semana 20).

Numeración del repositorio: la Semana 19 cierra con 325 pruebas, así
que la Semana 20 añade 326-345 (20 pruebas → 345 en total): plantillas
con marca y fila EJEMPLO (326-327), lectura de CSV y .xlsx con BOM y
delimitador detectado (328), límites de tamaño y filas (329), fechas
dd/mm/aaaa y aaaa-mm-dd (330), DNI de 7 u 8 dígitos enmascarado (331),
altas de productores idempotentes (332), errores con sugerencias e
informe sin datos personales (333), parcelas por productor+referencia
(334), campañas completas cerradas con sus 4 recursos (335), la atípica
que queda cosechada y marcada (336), filas malas sin dejar resto (337),
dry-run sin escribir (338), permisos y usuario (339), una EN_PROCESO y
--forzar (340), el registro de importaciones (341), un evento de
auditoría por ejecución (342), evaluar_todo una sola vez (343), el
dataset con origen y NULLs (344) y el recorrido completo con los CSV
de docs/ejemplos_importacion (345).
"""
import os
import tempfile
from datetime import date
from decimal import Decimal
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from auditoria.models import Auditoria
from campanias.models import Campania
from catalogos.models import Cultivo, Unidad, Variedad
from importaciones import servicio
from importaciones.models import Importacion
from importaciones.plantillas import (
    COLUMNA_MARCA,
    COLUMNAS,
    EJEMPLOS,
    VALOR_EJEMPLO,
    encabezados,
    generar_plantilla,
)
from productores.models import Comunidad, Distrito, Parcela, Productor, Provincia
from produccion.models import Produccion
from seguimiento.models import UsoRecurso
from usuarios.models import Rol, Usuario

EJEMPLOS_DIR = Path(settings.BASE_DIR).parent / "docs" / "ejemplos_importacion"

PRODUCTORES_CSV = (
    "marca,nombres,apellidos,dni,telefono,comunidad,referencia_ubicacion,"
    "fecha_incorporacion\n"
    "EJEMPLO,Ana,Ejemplo Lopez,88000001,999999999,Comunidad de prueba,"
    "Cerca del caserio,01/01/2020\n"
    ",Lucia,Prueba S20,77770001,911222333,Comunidad S20,Detras del almacen,"
    "15/03/2019\n"
    ",Mateo,Prueba S20,77770002,,Comunidad S20,Lado del rio,2019-06-01\n"
)

CAMPAÑAS_CSV = (
    "dni_productor,referencia_parcela,cultivo,variedad,periodo,fecha_inicio,"
    "fecha_fin_prevista,fecha_siembra,superficie_sembrada_ha,fecha_cosecha,"
    "superficie_cosechada_ha,produccion_kg,semilla_kg,fertilizante_kg,jornales,"
    "agua_m3,observaciones\n"
    "77770001,Terreno S20,Quinua S20,Blanca S20,2024-2025,01/11/2024,"
    "30/04/2025,10/11/2024,,20/04/2025,,2000,45,120,30,200,Fila de ejemplo\n"
)


class ImportacionBase(TestCase):
    """Suelo común: admin con el permiso, coordinador sin él, ubicación,
    cultivo con variedad y los recursos genéricos de la migración."""

    def setUp(self):
        self.admin = Usuario.objects.create_superuser(
            username="admin_s20", email="", password="AdminS20!",
            rol=Rol.objects.get(nombre="Administrador"),
        )
        self.coordinador = Usuario.objects.create_user(
            username="coord_s20", password="CoordS20!",
            rol=Rol.objects.get(nombre="Coordinador"),
        )
        provincia = Provincia.objects.create(nombre="Provincia S20")
        distrito = Distrito.objects.create(
            provincia=provincia, nombre="Distrito S20"
        )
        self.comunidad = Comunidad.objects.create(
            distrito=distrito, nombre="Comunidad S20"
        )
        self.cultivo = Cultivo.objects.create(nombre="Quinua S20")
        self.variedad = Variedad.objects.create(
            cultivo=self.cultivo, nombre="Blanca S20"
        )

    # ------------------------------------------------------------------ #
    # helpers
    # ------------------------------------------------------------------ #
    def _archivo(self, contenido, sufijo=".csv"):
        descriptor, ruta = tempfile.mkstemp(suffix=sufijo)
        with os.fdopen(descriptor, "w", encoding="utf-8-sig", newline="") as archivo:
            archivo.write(contenido)
        self.addCleanup(os.unlink, ruta)
        return ruta

    def _xlsx(self, filas, tipo="productores"):
        descriptor, ruta = tempfile.mkstemp(suffix=".xlsx")
        os.close(descriptor)
        libro = Workbook()
        hoja = libro.active
        hoja.append(list(encabezados(tipo)))
        for fila in filas:
            hoja.append(fila)
        libro.save(ruta)
        self.addCleanup(os.unlink, ruta)
        return ruta

    def _ejecutar(self, usuario="admin_s20", **opciones):
        opciones.setdefault("usuario", usuario)
        salida = StringIO()
        call_command("importar_historico", stdout=salida, **opciones)
        return salida.getvalue()

    def _productor(self, dni="77770001", nombres="Lucia", apellidos="Prueba S20"):
        return Productor.objects.create(
            nombres=nombres,
            apellidos=apellidos,
            dni=dni,
            comunidad=self.comunidad,
            creado_por=self.admin,
        )

    def _parcela(self, productor, referencia="Terreno S20", superficie="2.50"):
        return Parcela.objects.create(
            productor=productor,
            comunidad=self.comunidad,
            referencia_ubicacion=referencia,
            superficie_ha=superficie,
            altitud_msnm=3800,
            creado_por=self.admin,
        )


class PlantillaTests(ImportacionBase):
    """Pruebas 326-327: las plantillas .xlsx que genera el sistema."""

    def test_326_las_plantillas_llevan_encabezados_y_fila_ejemplo(self):
        for tipo in ("productores", "parcelas", "campanias"):
            descriptor, ruta = tempfile.mkstemp(suffix=f"_{tipo}.xlsx")
            os.close(descriptor)
            self.addCleanup(os.unlink, ruta)
            generar_plantilla(tipo, ruta)
            libro = load_workbook(ruta)
            try:
                hoja = libro.active
                self.assertEqual(hoja.title, tipo)
                primera = [celda.value for celda in hoja[1]]
                segunda = [celda.value for celda in hoja[2]]
            finally:
                libro.close()
            self.assertEqual(primera[0], COLUMNA_MARCA)
            self.assertEqual(tuple(primera[1:]), COLUMNAS[tipo])
            self.assertEqual(len(primera), len(COLUMNAS[tipo]) + 1)
            self.assertEqual(segunda[0], VALOR_EJEMPLO)
            self.assertEqual(tuple(segunda), EJEMPLOS[tipo])
            self.assertEqual(len(EJEMPLOS[tipo]), len(encabezados(tipo)))

    def test_327_el_comando_escribe_la_plantilla_pedida(self):
        ruta = os.path.join(
            tempfile.gettempdir(), "plantilla_comando_s20.xlsx"
        )
        self.addCleanup(lambda: os.path.exists(ruta) and os.unlink(ruta))
        salida = StringIO()
        call_command(
            "importar_historico", "--plantilla", "campanias", "--salida", ruta,
            stdout=salida,
        )
        self.assertIn("Plantilla escrita", salida.getvalue())
        self.assertTrue(os.path.exists(ruta))
        with self.assertRaisesMessage(CommandError, "--salida"):
            call_command("importar_historico", "--plantilla", "campanias")


class LecturaTests(ImportacionBase):
    """Pruebas 328-329: cómo se leen los archivos y sus límites."""

    def test_328_lee_csv_con_bom_punto_y_coma_y_xlsx_sin_fila_ejemplo(self):
        csv_pyc = self._archivo(
            "nombres;apellidos;dni;telefono;comunidad;referencia_ubicacion;"
            "fecha_incorporacion\n"
            "Irene;Punto Y Coma;77771001;;Comunidad S20;Ruta antigua;01/01/2020\n"
        )
        filas = servicio.leer_archivo("productores", csv_pyc)
        self.assertEqual(len(filas), 1)
        numero, fila = filas[0]
        self.assertEqual(numero, 2)
        self.assertEqual(fila["dni"], "77771001")
        self.assertEqual(fila["nombres"], "Irene")

        xlsx = self._xlsx(
            [
                list(EJEMPLOS["productores"]),
                ["", "Rosa", "Del Xlsx", "77771002", "", "Comunidad S20",
                 "Detras del almacen", "05/05/2020"],
            ]
        )
        filas = servicio.leer_archivo("productores", xlsx)
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][1]["dni"], "77771002")
        self.assertEqual(filas[0][1]["fecha_incorporacion"], "05/05/2020")

        sin_columna = self._archivo("nombres,apellidos\nLucia,Prueba\n")
        with self.assertRaisesMessage(servicio.ErrorArchivo, "Faltan columnas"):
            servicio.leer_archivo("productores", sin_columna)

    def test_329_rechaza_archivos_grandes_y_de_mas_de_5000_filas(self):
        grande = self._archivo(
            "nombres,apellidos,dni,telefono,comunidad,referencia_ubicacion,"
            "fecha_incorporacion\n" + "x" * (servicio.LIMITE_BYTES + 10)
        )
        with self.assertRaisesMessage(servicio.ErrorArchivo, "5 MB"):
            servicio.leer_archivo("productores", grande)

        filas = [
            "nombres,apellidos,dni,telefono,comunidad,referencia_ubicacion,"
            "fecha_incorporacion"
        ]
        filas += [
            f"Persona{i},Limite,7779{i:04d},,Comunidad S20,Ref{i},01/01/2020"
            for i in range(servicio.LIMITE_FILAS + 1)
        ]
        muchas = self._archivo("\n".join(filas) + "\n")
        with self.assertRaisesMessage(servicio.ErrorArchivo, "el límite es 5000"):
            servicio.leer_archivo("productores", muchas)


class NormalizacionTests(ImportacionBase):
    """Pruebas 330-331: fechas y DNI."""

    def test_330_solo_admite_dd_mm_aaaa_y_aaaa_mm_dd(self):
        self.assertEqual(servicio.parsear_fecha("15/03/2019", "f"), date(2019, 3, 15))
        self.assertEqual(servicio.parsear_fecha("2019-03-15", "f"), date(2019, 3, 15))
        with self.assertRaises(servicio.ErroresFila) as contexto:
            servicio.parsear_fecha("", "fecha_cosecha")
        self.assertEqual(contexto.exception.errores[0][0], "fecha_cosecha")
        with self.assertRaisesMessage(servicio.ErroresFila, "Fecha inválida"):
            servicio.parsear_fecha("31/02/2020", "f")
        with self.assertRaisesMessage(servicio.ErroresFila, "no reconocida"):
            servicio.parsear_fecha("15-03-2019", "f")
        with self.assertRaisesMessage(servicio.ErroresFila, "no reconocida"):
            servicio.parsear_fecha("1/2/2020", "f")

    def test_331_el_dni_de_siete_digitos_lleva_cero_y_los_mensajes_lo_enmascaran(self):
        self.assertEqual(servicio.normalizar_dni("7770001"), "07770001")
        self.assertEqual(servicio.normalizar_dni("07770001"), "07770001")
        self.assertEqual(servicio.enmascarar_dni("07770001"), "******01")
        self.assertNotIn("07770001", servicio.enmascarar_dni("07770001"))
        with self.assertRaisesMessage(servicio.ErroresFila, "7 u 8 dígitos"):
            servicio.normalizar_dni("77700")
        with self.assertRaisesMessage(servicio.ErroresFila, "7 u 8 dígitos"):
            servicio.normalizar_dni("1234567a")

        ruta = self._archivo(
            "nombres,apellidos,dni,telefono,comunidad,referencia_ubicacion,"
            "fecha_incorporacion\n"
            "Nadia,Siete Digitos,7777200,,Comunidad S20,Ref,15/03/2019\n"
        )
        salida = self._ejecutar(tipo="productores", archivo=ruta)
        self.assertRegex(salida, r"creadas:\s+1\b")
        self.assertTrue(Productor.objects.filter(dni="07777200").exists())


class ProductoresTests(ImportacionBase):
    """Pruebas 332-333: altas, duplicados y errores útiles."""

    def test_332_crea_productores_y_re_ejecutar_los_omite_sin_duplicar(self):
        ruta = self._archivo(PRODUCTORES_CSV)
        primera = self._ejecutar(tipo="productores", archivo=ruta)
        self.assertRegex(primera, r"creadas:\s+2\b")
        self.assertRegex(primera, r"ya existían:\s+0\b")
        self.assertRegex(primera, r"con error:\s+0\b")

        lucia = Productor.objects.get(dni="77770001")
        self.assertEqual(lucia.nombres, "Lucia")
        self.assertEqual(lucia.comunidad, self.comunidad)
        self.assertEqual(lucia.fecha_incorporacion, date(2019, 3, 15))
        self.assertEqual(lucia.creado_por, self.admin)
        sin_telefono = Productor.objects.get(dni="77770002")
        self.assertIn(sin_telefono.telefono, (None, ""))

        segunda = self._ejecutar(tipo="productores", archivo=ruta)
        self.assertRegex(segunda, r"creadas:\s+0\b")
        self.assertRegex(segunda, r"ya existían:\s+2\b")
        self.assertEqual(Productor.objects.count(), 2)

    def test_333_los_errores_sugieren_y_el_informe_no_descubre_dnis(self):
        ruta = self._archivo(
            "nombres,apellidos,dni,telefono,comunidad,referencia_ubicacion,"
            "fecha_incorporacion\n"
            "Eva,Comunidad Rara,77770003,,Comunidad S20 x,Calle 1,01/01/2020\n"
            "Pepe,Dni Corto,77700,,Comunidad S20,Otro lado,01/01/2020\n"
        )
        parcela_mala = self._archivo(
            "dni_productor,comunidad,referencia_ubicacion,superficie_ha,"
            "altitud_msnm,latitud,longitud,tipo_suelo,pendiente,"
            "disponibilidad_agua,sistema_riego\n"
            "77123456,Comunidad S20,Parce lejana,1.00,3800,,,,,,\n"
        )
        with tempfile.TemporaryDirectory() as carpeta:
            with patch.dict(os.environ, {"IMPORTACION_DIR": carpeta}):
                salida = self._ejecutar(tipo="productores", archivo=ruta)
                salida_parcelas = self._ejecutar(
                    tipo="parcelas", archivo=parcela_mala
                )
            self.assertRegex(salida, r"con error:\s+2\b")
            self.assertIn("Quizá quiso decir: Comunidad S20", salida)
            self.assertIn("7 u 8 dígitos", salida)
            self.assertEqual(Productor.objects.count(), 0)

            self.assertIn("******56", salida_parcelas)
            self.assertNotIn("77123456", salida_parcelas)
            informes = sorted(Path(carpeta).glob("informe_importacion_*.csv"))
            self.assertEqual(len(informes), 2)
            parcelas_informe = [p for p in informes if "parcelas" in p.name]
            self.assertEqual(len(parcelas_informe), 1)
            contenido = parcelas_informe[0].read_text(encoding="utf-8-sig")
            self.assertTrue(
                contenido.splitlines()[0].startswith("fila,columna,mensaje")
            )
            self.assertIn("dni_productor", contenido)
            self.assertNotIn("77123456", contenido)


class ParcelasTests(ImportacionBase):
    """Prueba 334: la clave es productor + referencia de ubicación."""

    def test_334_las_parcelas_se_clavan_por_productor_y_referencia(self):
        productor = self._productor()
        ruta = self._archivo(
            "dni_productor,comunidad,referencia_ubicacion,superficie_ha,"
            "altitud_msnm,latitud,longitud,tipo_suelo,pendiente,"
            "disponibilidad_agua,sistema_riego\n"
            "77770001,Comunidad S20,Terreno S20,2.50,3800,-15.500000,"
            "-70.100000,franco,suave,media,gravedad\n"
            "77770001,Comunidad S20,,1.20,3750,,,,,,\n"
        )
        primera = self._ejecutar(tipo="parcelas", archivo=ruta)
        self.assertRegex(primera, r"creadas:\s+1\b")
        self.assertRegex(primera, r"con error:\s+1\b")
        self.assertIn("referencia_ubicacion", primera)
        parcela = Parcela.objects.get(productor=productor)
        self.assertEqual(parcela.referencia_ubicacion, "Terreno S20")
        self.assertEqual(parcela.superficie_ha, Decimal("2.50"))
        self.assertEqual(parcela.tipo_suelo, "FRANCO")

        segunda = self._ejecutar(tipo="parcelas", archivo=ruta)
        self.assertRegex(segunda, r"ya existían:\s+1\b")
        self.assertEqual(Parcela.objects.count(), 1)


class CampaniasTests(ImportacionBase):
    """Pruebas 335-337: campañas completas, atípicas y atomicidad."""

    def setUp(self):
        super().setUp()
        self.productor = self._productor()
        self.parcela = self._parcela(self.productor)

    def test_335_importa_campania_completa_cerrada_con_sus_cuatro_recursos(self):
        ruta = self._archivo(CAMPAÑAS_CSV)
        salida = self._ejecutar(tipo="campanias", archivo=ruta)
        self.assertRegex(salida, r"creadas:\s+1\b")
        self.assertRegex(salida, r"con error:\s+0\b")

        campania = Campania.objects.get(periodo="2024-2025")
        self.assertEqual(campania.estado, Campania.ESTADO_CERRADA)
        self.assertEqual(campania.origen, Campania.ORIGEN_IMPORTADA)
        self.assertEqual(campania.creado_por, self.admin)
        self.assertEqual(campania.superficie_sembrada_ha, Decimal("2.50"))
        self.assertEqual(campania.fecha_siembra, date(2024, 11, 10))

        produccion = Produccion.objects.get(campania=campania)
        self.assertEqual(produccion.rendimiento_kg_ha, Decimal("800.00"))
        self.assertEqual(produccion.superficie_cosechada_ha, Decimal("2.50"))
        self.assertEqual(produccion.fecha_cosecha, date(2025, 4, 20))
        self.assertFalse(produccion.requiere_revision)
        self.assertTrue(produccion.incluir_en_entrenamiento)

        usos = UsoRecurso.objects.filter(campania=campania).order_by("id")
        self.assertEqual(usos.count(), 4)
        por_nombre = {uso.recurso.nombre: uso for uso in usos}
        self.assertEqual(set(por_nombre), set(servicio.RECURSOS_GENERICOS.values()))
        base_masa = Unidad.objects.get(tipo="MASA", es_base=True)
        base_labor = Unidad.objects.get(tipo="LABOR", es_base=True)
        base_volumen = Unidad.objects.get(tipo="VOLUMEN", es_base=True)
        esperado = {
            "semilla_kg": (Decimal("45"), base_masa),
            "fertilizante_kg": (Decimal("120"), base_masa),
            "jornales": (Decimal("30"), base_labor),
            "agua_m3": (Decimal("200"), base_volumen),
        }
        for clave, (cantidad, unidad) in esperado.items():
            uso = por_nombre[servicio.RECURSOS_GENERICOS[clave]]
            self.assertEqual(uso.cantidad, cantidad)
            self.assertEqual(uso.unidad, unidad)
            self.assertEqual(uso.cantidad_base, cantidad)
            self.assertEqual(uso.fecha, campania.fecha_siembra)

    def test_336_un_rendimiento_atipico_queda_cosechada_y_marcada(self):
        atipica = (
            "dni_productor,referencia_parcela,cultivo,variedad,periodo,"
            "fecha_inicio,fecha_fin_prevista,fecha_siembra,"
            "superficie_sembrada_ha,fecha_cosecha,superficie_cosechada_ha,"
            "produccion_kg,semilla_kg,fertilizante_kg,jornales,agua_m3,"
            "observaciones\n"
            "77770001,Terreno S20,Quinua S20,Blanca S20,2025-2026,01/11/2025,"
            "30/04/2026,10/11/2025,2.50,20/04/2026,,5,40,100,25,150,"
            "Rendimiento a proposito\n"
        )
        salida = self._ejecutar(tipo="campanias", archivo=self._archivo(atipica))
        self.assertRegex(salida, r"creadas:\s+1\b")

        campania = Campania.objects.get(periodo="2025-2026")
        self.assertEqual(campania.estado, Campania.ESTADO_COSECHADA)
        produccion = Produccion.objects.get(campania=campania)
        self.assertTrue(produccion.requiere_revision)
        self.assertIn("atípico", produccion.motivo_revision)
        self.assertEqual(produccion.rendimiento_kg_ha, Decimal("2.00"))

    def test_337_una_fila_mala_no_deja_nada_a_medias(self):
        with self.subTest("catálogo inexistente"):
            ruta = self._archivo(
                CAMPAÑAS_CSV
                + "77770001,Terreno S20,Quinua S20,Variedad Fantasma,2026-2027,"
                "01/11/2026,30/04/2027,10/11/2026,2.50,20/04/2027,,900,,,,,"
                "Fila mala\n"
            )
            salida = self._ejecutar(tipo="campanias", archivo=ruta)
            self.assertRegex(salida, r"creadas:\s+1\b")
            self.assertRegex(salida, r"con error:\s+1\b")
            self.assertFalse(Campania.objects.filter(periodo="2026-2027").exists())
            self.assertEqual(UsoRecurso.objects.count(), 4)
            self.assertEqual(Produccion.objects.count(), 1)

        with self.subTest("cosecha antes de la siembra"):
            ruta = self._archivo(
                CAMPAÑAS_CSV
                + "77770001,Terreno S20,Quinua S20,Blanca S20,2027-2028,"
                "01/11/2027,30/04/2028,10/11/2027,2.50,01/11/2026,,900,,,,,"
                "Cosecha imposible\n"
            )
            salida = self._ejecutar(tipo="campanias", archivo=ruta)
            self.assertRegex(salida, r"con error:\s+1\b")
            self.assertIn("anterior a la fecha de siembra", salida)
            self.assertFalse(Campania.objects.filter(periodo="2027-2028").exists())


class ComandoTests(ImportacionBase):
    """Pruebas 338-343: dry-run, permisos, EN_PROCESO, registro,
    auditoría y las alertas del cierre."""

    def test_338_el_dry_run_valida_y_no_escribe_nada(self):
        ruta = self._archivo(PRODUCTORES_CSV)
        salida = self._ejecutar(tipo="productores", archivo=ruta, dry_run=True)
        self.assertIn("dry-run", salida)
        self.assertRegex(salida, r"creadas:\s+2\b")
        self.assertEqual(Productor.objects.count(), 0)
        self.assertEqual(Importacion.objects.count(), 0)
        self.assertEqual(Auditoria.objects.filter(modulo="importaciones").count(), 0)

    def test_339_pide_permiso_de_importacion_y_un_usuario_activo(self):
        ruta = self._archivo(PRODUCTORES_CSV)
        with self.assertRaisesMessage(CommandError, "importacion.ejecutar"):
            self._ejecutar(usuario="coord_s20", tipo="productores", archivo=ruta)
        self.assertEqual(Productor.objects.count(), 0)

        with self.assertRaisesMessage(CommandError, "No existe el usuario"):
            self._ejecutar(usuario="nadie", tipo="productores", archivo=ruta)

        inactivo = Usuario.objects.create_user(
            username="baja_s20", password="BajaS20!",
            rol=Rol.objects.get(nombre="Administrador"), is_active=False,
        )
        with self.assertRaisesMessage(CommandError, "inactivo"):
            self._ejecutar(usuario=inactivo.username, tipo="productores", archivo=ruta)

        with self.assertRaisesMessage(CommandError, "Faltan argumentos"):
            call_command("importar_historico", tipo="productores")

    def test_340_solo_puede_haber_una_importacion_en_proceso_y_forzar_larga(self):
        previa = Importacion.objects.create(
            archivo="interrumpido.csv",
            tipo="PRODUCTORES",
            usuario=self.admin,
            filas_totales=10,
            estado="EN_PROCESO",
        )
        ruta = self._archivo(PRODUCTORES_CSV)
        with self.assertRaisesMessage(CommandError, "EN_PROCESO"):
            self._ejecutar(tipo="productores", archivo=ruta)

        salida = self._ejecutar(tipo="productores", archivo=ruta, forzar=True)
        self.assertIn("--forzar", salida)
        previa.refresh_from_db()
        self.assertEqual(previa.estado, "CON_ERRORES")
        self.assertIn("--forzar", previa.detalle_errores[0]["mensaje"])
        self.assertRegex(salida, r"creadas:\s+2\b")

    def test_341_el_registro_de_importaciones_guarda_totales_y_errores(self):
        self._parcela(self._productor())
        ruta = self._archivo(
            CAMPAÑAS_CSV
            + "77770001,Terreno S20,Quinua S20,Variedad Fantasma,2026-2027,"
            "01/11/2026,30/04/2027,10/11/2026,2.50,20/04/2027,2.50,900,,,,,"
            "Fila mala\n"
        )
        self._ejecutar(tipo="campanias", archivo=ruta)
        registro = Importacion.objects.get(tipo="CAMPANIAS")
        self.assertEqual(registro.estado, "CON_ERRORES")
        self.assertEqual(registro.filas_totales, 2)
        self.assertEqual(registro.filas_ok, 1)
        self.assertEqual(registro.filas_error, 1)
        self.assertEqual(len(registro.detalle_errores), 1)
        self.assertEqual(registro.detalle_errores[0]["fila"], 3)
        self.assertEqual(registro.usuario, self.admin)

        limpia = self._ejecutar(tipo="campanias", archivo=self._archivo(CAMPAÑAS_CSV))
        self.assertRegex(limpia, r"con error:\s+0\b")
        registro_limpio = Importacion.objects.filter(
            tipo="CAMPANIAS", estado="COMPLETADA"
        ).get()
        self.assertEqual(registro_limpio.filas_error, 0)
        self.assertIsNone(registro_limpio.detalle_errores)

    def test_342_cada_ejecucion_deja_un_evento_de_auditoria(self):
        base = Auditoria.objects.filter(modulo="importaciones").count()
        ruta = self._archivo(PRODUCTORES_CSV)
        self._ejecutar(tipo="productores", archivo=ruta)
        self.assertEqual(
            Auditoria.objects.filter(modulo="importaciones").count(), base + 1
        )
        evento = Auditoria.objects.filter(
            modulo="importaciones", tabla="importaciones"
        ).latest("id")
        self.assertEqual(evento.accion, "CREAR")
        self.assertEqual(evento.usuario, self.admin)
        self.assertEqual(evento.detalle["tipo"], "productores")
        self.assertEqual(evento.detalle["filas_ok"], 2)

    def test_343_las_alertas_se_evaluan_una_vez_y_nunca_en_dry_run(self):
        ruta = self._archivo(PRODUCTORES_CSV)
        with patch(
            "importaciones.management.commands.importar_historico.evaluar_todo",
            return_value={"creadas": 0, "cerradas": 0},
        ) as parche:
            self._ejecutar(tipo="productores", archivo=ruta)
            parche.assert_called_once_with()

        with patch(
            "importaciones.management.commands.importar_historico.evaluar_todo",
            return_value={"creadas": 0, "cerradas": 0},
        ) as parche:
            self._ejecutar(tipo="productores", archivo=ruta, dry_run=True)
            parche.assert_not_called()


class DatasetTests(ImportacionBase):
    """Pruebas 344-345: el dataset de entrenamiento y el recorrido
    completo con los CSV de ejemplo del repositorio."""

    def setUp(self):
        super().setUp()
        self.productor = self._productor()
        self.parcela = self._parcela(self.productor)

    def _fila_dataset(self, campania):
        from django.db import connection

        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT origen, n_heladas, n_incidencias, severidad_max, "
                "n_aplicaciones, dias_ventana_clima, rendimiento_kg_ha "
                "FROM vw_dataset_entrenamiento WHERE campania_id = %s",
                [campania.pk],
            )
            return cursor.fetchone()

    def test_344_las_campanias_importadas_al_dataset_con_origen_y_nulls(self):
        self._ejecutar(tipo="campanias", archivo=self._archivo(CAMPAÑAS_CSV))
        campania = Campania.objects.get(periodo="2024-2025")
        fila = self._fila_dataset(campania)
        self.assertIsNotNone(fila)
        origen, heladas, incidencias, severidad, aplicaciones, ventana, rendimiento = (
            fila
        )
        self.assertEqual(origen, "IMPORTADA")
        self.assertIsNone(heladas)
        self.assertIsNone(incidencias)
        self.assertIsNone(severidad)
        self.assertIsNone(aplicaciones)
        self.assertEqual(ventana, 162)
        self.assertEqual(rendimiento, Decimal("800.00"))

        atipica = (
            "dni_productor,referencia_parcela,cultivo,variedad,periodo,"
            "fecha_inicio,fecha_fin_prevista,fecha_siembra,"
            "superficie_sembrada_ha,fecha_cosecha,superficie_cosechada_ha,"
            "produccion_kg,semilla_kg,fertilizante_kg,jornales,agua_m3,"
            "observaciones\n"
            "77770001,Terreno S20,Quinua S20,Blanca S20,2025-2026,01/11/2025,"
            "30/04/2026,10/11/2025,2.50,20/04/2026,,5,,,,,Atipica\n"
        )
        self._ejecutar(tipo="campanias", archivo=self._archivo(atipica))
        self.assertIsNone(self._fila_dataset(Campania.objects.get(periodo="2025-2026")))

    def test_345_el_recorrido_completo_con_los_csv_de_ejemplo(self):
        self.assertTrue(
            EJEMPLOS_DIR.is_dir(), f"Falta la carpeta {EJEMPLOS_DIR}"
        )
        call_command("cargar_ubicacion", stdout=StringIO())
        self.assertTrue(
            Comunidad.objects.filter(nombre="Comunidad Demo N1-1").exists()
        )
        quinua, _ = Cultivo.objects.get_or_create(nombre="Quinua")
        Variedad.objects.get_or_create(cultivo=quinua, nombre="Blanca")

        with tempfile.TemporaryDirectory() as carpeta:
            with patch.dict(os.environ, {"IMPORTACION_DIR": carpeta}):
                for tipo in ("productores", "parcelas", "campanias"):
                    ruta = EJEMPLOS_DIR / f"{tipo}.csv"
                    salida = self._ejecutar(tipo=tipo, archivo=str(ruta))
                    self.assertRegex(
                        salida, r"con error:\s+0\b", f"{tipo}: {salida}"
                    )

        self.assertEqual(
            Productor.objects.filter(dni__startswith="7788000").count(), 3
        )
        self.assertEqual(
            Parcela.objects.filter(productor__dni__startswith="7788000").count(), 3
        )
        campanias = Campania.objects.filter(origen="IMPORTADA")
        self.assertEqual(campanias.count(), 3)
        self.assertEqual(
            campanias.filter(estado=Campania.ESTADO_CERRADA).count(), 3
        )
        cierre = self._ejecutar(
            tipo="campanias", archivo=str(EJEMPLOS_DIR / "campanias.csv")
        )
        self.assertRegex(cierre, r"ya existían:\s+3\b")

        from django.db import connection

        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COUNT(*) FROM vw_dataset_entrenamiento "
                "WHERE origen = 'IMPORTADA'"
            )
            self.assertEqual(cursor.fetchone()[0], 3)
